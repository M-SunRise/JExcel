import os
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from typing import List, Dict, Any, Optional


class ExcelExporter:
    HEADER_FILL = PatternFill(start_color="1E90FF", end_color="1E90FF", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    CELL_ALIGNMENT = Alignment(horizontal="left", vertical="center")
    BORDER = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    def export(self, flat_data: List[Dict], output_path: Optional[str] = None,
               auto_width: bool = True) -> Dict[str, Any]:
        try:
            if not flat_data:
                return {"success": False, "error": "没有可导出的数据"}

            headers = list(flat_data[0].keys())
            wb = Workbook()
            ws = wb.active
            ws.title = "Data"

            for c_idx, col_name in enumerate(headers, 1):
                cell = ws.cell(row=1, column=c_idx, value=col_name)
                cell.fill = self.HEADER_FILL
                cell.font = self.HEADER_FONT
                cell.alignment = self.CELL_ALIGNMENT
                cell.border = self.BORDER

            for r_idx, record in enumerate(flat_data, 2):
                for c_idx, col_name in enumerate(headers, 1):
                    value = record.get(col_name, '')
                    if value is None:
                        value = ''
                    cell = ws.cell(row=r_idx, column=c_idx, value=value)
                    cell.alignment = self.CELL_ALIGNMENT
                    cell.border = self.BORDER

            if auto_width:
                for col_idx in range(1, len(headers) + 1):
                    max_length = 0
                    col_letter = ws.cell(row=1, column=col_idx).column_letter
                    for cell in ws[col_letter]:
                        try:
                            cell_len = len(str(cell.value or ''))
                            if cell_len > max_length:
                                max_length = cell_len
                        except:
                            pass
                    ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

            ws.freeze_panes = "A2"

            if output_path:
                wb.save(output_path)
                return {
                    "success": True,
                    "rows": len(flat_data),
                    "columns": len(headers),
                    "path": output_path
                }

            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return {
                "success": True,
                "rows": len(flat_data),
                "columns": len(headers),
                "buffer": buffer
            }

        except Exception as e:
            return {"success": False, "error": str(e)}
